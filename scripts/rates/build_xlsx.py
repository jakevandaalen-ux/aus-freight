import json,sys,statistics as st,collections
sys.path.insert(0,'/home/claude/au/newrates')
from model import EST,INFO,STS,WHST,rows as MROWS,P,WH,over
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
from openpyxl.utils import get_column_letter as CL
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment
A='/home/claude/au/'
W=json.load(open(A+'webrates.json'));R=W['R'];E=W['E']
ORDER=json.load(open('order.json'))
SZ=['Small','Medium','Large','XL','XXL'];ZN=['Metro','Metro Outer','Regional']
# add missing columns
ORDER['TAS'].append('SA');ORDER['WA'].append('TAS')
# default priority ratio by size
rat=collections.defaultdict(list)
for d in R:
    for f in R[d]:
        for s in range(5):
            for z in range(3):
                a=R[d][f]['standard'][s][z];b=R[d][f]['priority'][s][z]
                if a and b: rat[s].append(b/a)
DR=[round(min(1.35,max(1.2,st.median(rat[s]))),2) for s in range(5)]
F='Arial';BLUE=Font(name=F,color='0000FF');BOLD=Font(name=F,bold=True);NORM=Font(name=F);GREEN=Font(name=F,color='008000')
H1=Font(name=F,bold=True,size=14);HDR=PatternFill('solid',fgColor='DDE6EE');YEL=PatternFill('solid',fgColor='FFFF00');GREY=Font(name=F,color='666666',italic=True)
thin=Side(style='thin',color='BBBBBB');BOX=Border(top=thin,bottom=thin,left=thin,right=thin)
MONEY='$#,##0.00';MONEY0='$#,##0'
wb=Workbook()
# ---------- Read me / settings
rm=wb.active;rm.title='Read me'
rm['A1']='RRG Website Shipping · Australia · proposed rates';rm['A1'].font=H1
rm['A2']='Prepared 20 October 2026 from Roof Racks Galore order and carrier data. Same layout as "RRG Website Shipping OCT 2025": one tab per delivery state, a table for each warehouse the order ships from, Standard and Priority, Metro / Metro Outer / Regional.';rm['A2'].font=NORM
rm['A4']='Settings (change the yellow cells; every price in the workbook updates)';rm['A4'].font=BOLD
sets=[('Freight margin on top of expected cost',0.0,'0.0%','0% = break-even: on average the Standard charge covers what the carrier charges us for that size, zone and warehouse. 10% adds 10% on top of expected cost.'),
      ('Minimum Priority uplift over Standard ($)',5.0,MONEY,'Priority is at least this much more than Standard.'),
      ('Round prices up to the next $x.50 up to ($)',30.0,MONEY,'Below this, prices end in .50 (e.g. $12.50). Above it they round up to the next $5.')]
for i,(k,v,fmt,n) in enumerate(sets):
    r=5+i;rm.cell(r,1,k).font=NORM;c=rm.cell(r,2,v);c.font=BLUE;c.fill=YEL;c.number_format=fmt;rm.cell(r,3,n).font=GREY
rm['A9']='Priority ÷ Standard ratio by size';rm['A9'].font=BOLD
for s in range(5):
    rm.cell(10+s,1,SZ[s]).font=NORM;c=rm.cell(10+s,2,DR[s]);c.font=BLUE;c.number_format='0.00"x"'
rm.cell(10,3,'Median of Priority ÷ Standard across the 2025 sheet, kept between 1.20 and 1.35 (Large was 2.1 and XL 1.0 in 2025 because those Standard prices were far below cost). Applied to every cell. Change these if you want a different Priority premium.').font=GREY
notes=['How the prices are set',
 '1. Each order is sized the way the website does it: every product gets a 1–5 score from its longest side, cubic weight (250 kg/m³) and actual weight, and the order takes its biggest product\'s size. See the Freight SIZE RULES tab.',
 '2. Expected cost = what carriers actually charged us (invoiced cost, or the carrier estimate where not yet billed, plus 10% GST) on 5,882 matched orders, June–October 2026, that went as one invoice in one shipment.',
 '3. Costs are averaged after dropping the top and bottom 10% of orders in each group, so one-off surcharges don\'t set the price.',
 '4. Each size × zone gets a national figure, separately for same-state and interstate deliveries. Each warehouse-to-state lane then gets a factor from its own orders (for example WA to VIC costs about 25% more than the interstate average). Lanes with few orders lean on the national figure; lanes with no orders use distance.',
 '5. Where a lane has its own orders for a size and zone, those orders are blended in. Prices never go down from a smaller size to a larger one, or from Metro to Outer to Regional.',
 '6. Standard = expected cost × (1 + freight margin), rounded up. Priority = Standard × the Priority ratio for its size (above), at least the minimum uplift, rounded up. We don\'t yet have enough priority shipments to cost Priority separately.',
 '7. Delivery time frames are copied from the 2025 sheet unchanged.',
 '',
 'What else changed from the 2025 sheet',
 '• The TAS tab now has a From SA table and the WA tab a From TAS table (both were missing).',
 '• NT Priority XL and XXL now have prices from every warehouse (most were blank).',
 '• Postcode zones tab: every Australia Post postcode with its website zone name. 97 NSW Outer Metro postcodes (Central Coast, Newcastle, Wollongong, Blue Mountains, Coffs Harbour, Albury, Griffith) were missing from the 2025 Auspost template and are added, marked "Added".',
 '• Freight SIZE RULES tab now matches how the website actually scores products (the 2025 tab listed different length limits and no cubic weight).',
 '',
 'Things to know',
 '• These are break-even prices on today\'s carrier mix. Medium is the biggest change: most Medium orders are 1.05–1.8 m long (roof bars), which Australia Post can\'t carry, so they go StarTrack at $40–55. See the Summary tab for an option that splits Medium at 1.05 m.',
 '• Large, XL and XXL have far fewer orders than Small and Medium, especially interstate Regional, so those prices are less certain. The Cost model tab shows how many orders sit behind every price.',
 '• Carrier rate cards weren\'t used directly: they need each carrier\'s postcode zones, fuel levy and surcharges, which weren\'t in the exports. Costs are what we were actually charged.',
 '• The Impact tab re-prices the same 5,882 orders under these rates so you can see the effect before changing the website.']
for i,n in enumerate(notes):
    c=rm.cell(17+i,1,n);c.font=BOLD if n and not n[0] in '0123456789•' and i in (0,9,15) else NORM
rm.column_dimensions['A'].width=58;rm.column_dimensions['B'].width=12;rm.column_dimensions['C'].width=110
for r in range(17,17+len(notes)): rm.cell(r,1).alignment=Alignment(wrap_text=False)
MARGIN="'Read me'!$B$5";MINUP="'Read me'!$B$6";RND="'Read me'!$B$7"
def rnd(x): return f'IF({x}<={RND},CEILING({x}+0.5,1)-0.5,CEILING({x},5))'
# ---------- Cost model
cm=wb.create_sheet('Cost model')
hdr=['Key','Delivering to','Ships from','Size','Zone','Expected cost incl. GST','Orders in this cell','Orders on this lane','Lane factor','2025 Standard','2025 Priority','Priority ratio','New Standard','New Priority','Standard change']
for j,h in enumerate(hdr,1):
    c=cm.cell(1,j,h);c.font=BOLD;c.fill=HDR;c.alignment=Alignment(wrap_text=True,vertical='top')
r=2;ROW={}
for dst in STS:
    for fr in ORDER[dst]:
        for s in range(5):
            for z in range(3):
                n,nl,lf=INFO[(dst,fr,s,z)];e=EST[(dst,fr,s,z)]
                os_=R.get(dst,{}).get(fr,{}).get('standard',[[None]*3]*5)[s][z] if fr in R.get(dst,{}) else None
                op=R[dst][fr]['priority'][s][z] if fr in R.get(dst,{}) else None
                key=f'{dst}|{fr}|{SZ[s]}|{ZN[z]}'
                vals=[key,dst,fr,SZ[s],ZN[z],round(e,2),n,nl,round(lf,2),os_,op]
                for j,v in enumerate(vals,1):
                    c=cm.cell(r,j,v);c.font=NORM
                cm.cell(r,6).number_format=MONEY;cm.cell(r,10).number_format=MONEY;cm.cell(r,11).number_format=MONEY;cm.cell(r,9).number_format='0.00'
                cm.cell(r,12,f"=INDEX('Read me'!$B$10:$B$14,{s+1})").number_format='0.00'
                cm.cell(r,13,'='+rnd(f'F{r}*(1+{MARGIN})')).number_format=MONEY
                cm.cell(r,14,'='+rnd(f'MAX(M{r}*L{r},M{r}+{MINUP})')).number_format=MONEY
                cm.cell(r,15,f'=IF(ISNUMBER(J{r}),M{r}-J{r},"new")').number_format='+$#,##0.00;-$#,##0.00;-'
                for j in range(12,16): cm.cell(r,j).font=NORM
                ROW[(dst,fr,s,z)]=r;r+=1
LAST=r-1
for j,w in enumerate([24,10,10,8,12,12,10,10,9,11,11,9,12,12,11],1): cm.column_dimensions[CL(j)].width=w
cm.freeze_panes='B2';cm.auto_filter.ref=f'A1:O{LAST}'
cm.cell(1,6).comment=Comment('Trimmed average (top and bottom 10% dropped) of what carriers charged us incl. GST, for this size and zone, adjusted for the warehouse-to-state lane. See Read me.','Claude')
# ---------- State tabs
for dst in STS:
    ws=wb.create_sheet(dst)
    fr=ORDER[dst];ncol=len(fr)
    title=f'Freight Charges for {dst} Orders'
    ws.cell(1,1,title).font=H1
    def block(top,label,kind):
        ws.cell(top,1,label).font=BOLD
        for k,f in enumerate(fr):
            c0=2+k*4;ws.cell(top,c0,f'From {f}').font=BOLD
            for z in range(3):
                c=ws.cell(top+1,c0+z,ZN[z]);c.font=BOLD;c.fill=HDR;c.alignment=Alignment(horizontal='center')
        ws.cell(top+1,1,'Freight Size').font=BOLD;ws.cell(top+1,1).fill=HDR
        for s in range(5):
            ws.cell(top+2+s,1,SZ[s]).font=BOLD
            for k,f in enumerate(fr):
                for z in range(3):
                    c=ws.cell(top+2+s,2+k*4+z)
                    if kind in('standard','priority'):
                        rr=ROW[(dst,f,s,z)];col='M' if kind=='standard' else 'N'
                        c.value=f"='Cost model'!{col}{rr}";c.number_format=MONEY;c.font=GREEN
                    else:
                        svc=kind.split('-')[1]
                        src=f if f in E[dst] else ('VIC' if dst=='TAS' else 'VIC')
                        v=E[dst][src][svc][s][z];c.value=(v or '').strip().replace(' days',' Days').replace(' day',' Day') if v else None;c.font=NORM
                        if f not in E[dst]: c.comment=Comment(f'No {dst} from {f} times in the 2025 sheet; copied from VIC. Check.','Claude')
                    c.border=BOX;c.alignment=Alignment(horizontal='center')
    block(3,'Standard','standard');block(11,'Priority','priority')
    ws.cell(19,1,f'Delivery time frames for {dst} Orders').font=H1
    block(21,'Standard','eta-standard');block(29,'Priority','eta-priority')
    ws.column_dimensions['A'].width=13
    for k in range(ncol):
        for z in range(3): ws.column_dimensions[CL(2+k*4+z)].width=11
        ws.column_dimensions[CL(5+k*4)].width=2
    ws.freeze_panes='B3'
# ---------- Impact
im=wb.create_sheet('Impact')
ih=['Delivering to','Ships from','Zone','Size','Service charged','Same state','We charged','Cost us','2025 sheet price','New price','New − cost','Charged − cost']
for j,h in enumerate(ih,1):
    c=im.cell(1,j,h);c.font=BOLD;c.fill=HDR;c.alignment=Alignment(wrap_text=True,vertical='top')
M=json.load(open(A+'matched.json'))
r=2
for x in M:
    pi,wi,ci,ch,co=x[0],x[1],x[2],x[3],x[4];sz=x[16]
    if x[8]!=1 or x[9]!=1 or not sz or wi<0 or co<=0: continue
    z=P[pi][7]
    if z<0 or z>2: continue
    dst=over.get(pi,P[pi][3]);fr=WH[wi][0]
    if (dst,fr,sz-1,z) not in ROW: continue
    svc='Priority' if x[19]=='p' else 'Standard'
    old=x[18] if x[19]=='p' else x[17]
    vals=[dst,fr,ZN[z],SZ[sz-1],svc,'Yes' if dst==fr else 'No',round(ch,2),round(co*1.1,2),old]
    for j,v in enumerate(vals,1): im.cell(r,j,v).font=NORM
    rr=ROW[(dst,fr,sz-1,z)]
    im.cell(r,10,f"=IF(E{r}=\"Priority\",'Cost model'!N{rr},'Cost model'!M{rr})").font=GREEN
    im.cell(r,11,f'=J{r}-H{r}').font=NORM;im.cell(r,12,f'=G{r}-H{r}').font=NORM
    for j in (7,8,9,10,11,12): im.cell(r,j).number_format=MONEY
    r+=1
ILAST=r-1
for j,w in enumerate([10,10,12,8,10,8,11,11,11,11,11,11],1): im.column_dimensions[CL(j)].width=w
im.freeze_panes='A2';im.auto_filter.ref=f'A1:L{ILAST}'
# ---------- Summary
sm=wb.create_sheet('Summary',1)
sm['A1']='Summary';sm['A1'].font=H1
sm['A2']=f'Re-pricing the {ILAST-1:,} matched single-shipment orders (Impact tab) under the new rates. Updates when you change the settings on Read me.';sm['A2'].font=GREY
heads=['','Orders','We charged','Cost us','Charged − cost','New rates would charge','New − cost']
for j,h in enumerate(heads,1): c=sm.cell(4,j,h);c.font=BOLD;c.fill=HDR;c.alignment=Alignment(wrap_text=True)
I=f"Impact!$A$2:$L${ILAST}"
def col(c): return f"Impact!${c}$2:${c}${ILAST}"
r=5;sm.cell(r,1,'All orders').font=BOLD
sm.cell(r,2,f'=COUNTA({col("A")})');sm.cell(r,3,f'=SUM({col("G")})');sm.cell(r,4,f'=SUM({col("H")})');sm.cell(r,5,f'=C{r}-D{r}');sm.cell(r,6,f'=SUM({col("J")})');sm.cell(r,7,f'=F{r}-D{r}')
r=7;sm.cell(r,1,'By size').font=BOLD;r+=1
for s in range(5):
    sm.cell(r,1,SZ[s]).font=NORM;cond=f'{col("D")},"{SZ[s]}"'
    sm.cell(r,2,f'=COUNTIF({cond})');sm.cell(r,3,f'=SUMIF({cond},{col("G")})');sm.cell(r,4,f'=SUMIF({cond},{col("H")})');sm.cell(r,5,f'=C{r}-D{r}');sm.cell(r,6,f'=SUMIF({cond},{col("J")})');sm.cell(r,7,f'=F{r}-D{r}');r+=1
r+=1;sm.cell(r,1,'By zone').font=BOLD;r+=1
for z in range(3):
    sm.cell(r,1,ZN[z]).font=NORM;cond=f'{col("C")},"{ZN[z]}"'
    sm.cell(r,2,f'=COUNTIF({cond})');sm.cell(r,3,f'=SUMIF({cond},{col("G")})');sm.cell(r,4,f'=SUMIF({cond},{col("H")})');sm.cell(r,5,f'=C{r}-D{r}');sm.cell(r,6,f'=SUMIF({cond},{col("J")})');sm.cell(r,7,f'=F{r}-D{r}');r+=1
r+=1;sm.cell(r,1,'By warehouse').font=BOLD;r+=1
for w in ['QLD','NSW','VIC','SA','WA','TAS','ACT']:
    sm.cell(r,1,f'{w} · {WHST[w][1]}').font=NORM;cond=f'{col("B")},"{w}"'
    sm.cell(r,2,f'=COUNTIF({cond})');sm.cell(r,3,f'=SUMIF({cond},{col("G")})');sm.cell(r,4,f'=SUMIF({cond},{col("H")})');sm.cell(r,5,f'=C{r}-D{r}');sm.cell(r,6,f'=SUMIF({cond},{col("J")})');sm.cell(r,7,f'=F{r}-D{r}');r+=1
for rr in range(5,r):
    for j in range(3,8): sm.cell(rr,j).number_format='$#,##0;-$#,##0;-'
    for j in range(2,8): sm.cell(rr,j).font=NORM
SUMEND=r
# average price table for same-state metro etc
r+=1;sm.cell(r,1,'Typical price per order, Standard (median of each tab\'s same-state table)').font=BOLD;r+=1
for j,h in enumerate(['Size','Metro 2025','Metro new','Outer 2025','Outer new','Regional 2025','Regional new'],1): c=sm.cell(r,j,h);c.font=BOLD;c.fill=HDR
r+=1
for s in range(5):
    sm.cell(r,1,SZ[s]).font=NORM
    for z in range(3):
        o=[R[d][d]['standard'][s][z] for d in STS if d in R[d] and R[d][d]['standard'][s][z]]
        sm.cell(r,2+z*2,st.median(o) if o else None).number_format=MONEY
        cells=','.join(f"'Cost model'!M{ROW[(d,d,s,z)]}" for d in STS if (d,d,s,z) in ROW)
        sm.cell(r,3+z*2,f'=MEDIAN({cells})').number_format=MONEY
    r+=1
# medium split option
r+=1;sm.cell(r,1,'Option: split Medium at 1.05 m (what Medium orders actually cost us, incl. GST, trimmed average)').font=BOLD;r+=1
for j,h in enumerate(['','Orders','Metro','Outer metro','Regional'],1): c=sm.cell(r,j,h);c.font=BOLD;c.fill=HDR
r+=1
def tm(v):
    v=sorted(v);k=int(len(v)*0.1);v=v[k:len(v)-k] or v;return sum(v)/len(v) if v else None
for loc in (True,False):
    for lb,lab in ((False,'up to 1.05 m'),(True,'over 1.05 m')):
        L=[x for x in M if x[8]==1 and x[9]==1 and x[16]==2 and x[1]>=0 and x[4]>0 and 0<=P[x[0]][7]<=2 and ((over.get(x[0],P[x[0]][3])==WH[x[1]][0])==loc) and ((x[21]>1.05)==lb)]
        sm.cell(r,1,f'Medium {lab}, {"same state" if loc else "interstate"}').font=NORM;sm.cell(r,2,len(L)).font=NORM
        for z in range(3):
            v=[x[4]*1.1 for x in L if P[x[0]][7]==z];c=sm.cell(r,3+z,round(tm(v),2) if v else None);c.number_format=MONEY;c.font=NORM
        r+=1
sm.cell(r,1,'Most Medium orders under 1.05 m can go Australia Post. Splitting Medium would let the website keep a low price for those and charge the long ones what StarTrack costs. It needs a 6th size on the website, so it isn\'t applied in the state tabs.').font=GREY
for j,w in enumerate([44,10,14,14,14,18,14],1): sm.column_dimensions[CL(j)].width=w
# ---------- Freight size rules
fs=wb.create_sheet('Freight SIZE RULES')
fs['A1']='Freight size rules (how the website scores each product)';fs['A1'].font=H1
rules=[('Score','Size','Longest side','Cubic weight (L×W×H×250)','Actual weight'),
 (1,'Small','under 0.5 m','under 10 kg','under 5 kg'),(2,'Medium','0.5 to under 1.8 m','10 to under 25 kg','5 to under 15 kg'),
 (3,'Large','1.8 to under 3.5 m','25 to under 100 kg','15 to under 22 kg'),(4,'XL','3.5 to under 5 m','100 to under 180 kg','22 to under 35 kg'),(5,'XXL','5 m or more','180 kg or more','35 kg or more')]
for i,rw in enumerate(rules):
    for j,v in enumerate(rw,1):
        c=fs.cell(3+i,j,v);c.font=BOLD if i==0 else NORM
        if i==0:c.fill=HDR
fs['A10']='A product\'s size is the highest of its three scores. An order takes the size of its biggest product. Products with no dimensions or weight score 1 (Small) on that measure; 15,773 products in the 20 Oct 2026 export have neither, so the website treats them as Small.';fs['A10'].font=NORM
fs['A11']='The 2025 sheet\'s version of this tab listed Large 1.8–2.5 m, XL 2.5–3.0 m and XXL 3.0 m+, and no cubic weight. Those aren\'t the rules the website uses.';fs['A11'].font=GREY
for j,w in enumerate([8,10,22,28,20],1): fs.column_dimensions[CL(j)].width=w
# ---------- Source hierarchy
sh=wb.create_sheet('Source Hierarchy')
sh['A1']='Which warehouse the website ships from first (unchanged from the 2025 sheet; NT from the awaiting-approval version)';sh['A1'].font=BOLD
for j,h in enumerate(['Delivering to','1st','2nd','3rd','4th','5th'],1): c=sh.cell(3,j,h);c.font=BOLD;c.fill=HDR
for i,d in enumerate(STS):
    sh.cell(4+i,1,d).font=NORM
    for k,f in enumerate(W['pref'][d]): sh.cell(4+i,2+k,f'{f} · {WHST[f][1]}').font=NORM
for j in range(1,7): sh.column_dimensions[CL(j)].width=18
# ---------- Postcode zones
pz=wb.create_sheet('Postcode zones')
web=json.load(open('/home/claude/au/rates/web_rates.json'))['webname']
for j,h in enumerate(['Postcode','State tab','Website zone name','Zone','Australia Post group','Status','Suburbs'],1): c=pz.cell(1,j,h);c.font=BOLD;c.fill=HDR
AP=['Capital City','Major Area','Rural','Remote Islands & Locations']
WN={0:'Metro',1:'Outer Metro',2:'Rural'}
r=2
for i,p in enumerate(P):
    if p[7]<0: continue
    pc=p[0];dst=over.get(i,p[3])
    if pc in web: name=web[pc];status='Same as 2025'
    else:
        name=f'{dst} - {WN[min(p[7],2)]}' if not (dst in('ACT','NT','TAS') and p[7]==0) else dst
        status='Added'
    for j,v in enumerate([pc,dst,name,['Metro','Metro Outer','Regional','Regional'][p[7]],AP[p[7]],status,p[5][:80]],1): pz.cell(r,j,v).font=NORM
    r+=1
for j,w in enumerate([9,9,24,12,24,13,60],1): pz.column_dimensions[CL(j)].width=w
pz.freeze_panes='A2';pz.auto_filter.ref=f'A1:G{r-1}'
# tab order
order_tabs=['Read me','Summary']+STS+['Cost model','Impact','Freight SIZE RULES','Source Hierarchy','Postcode zones']
wb._sheets=[wb[n] for n in order_tabs]
out='/mnt/user-data/outputs/RRG Website Shipping AU - proposed rates OCT 2026.xlsx'
wb.save(out);print(out,LAST,ILAST)
